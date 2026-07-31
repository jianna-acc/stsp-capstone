# File: /backend/app/services/file_extraction.py
# Purpose: Extracts readable text from PDF, TXT, PPTX, XLSX,
# and legacy XLS study files and divides content into chunks.

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from pptx import Presentation
from pptx.exc import PackageNotFoundError, PythonPptxError
from xlrd.xldate import XLDateError, xldate_as_datetime


PDF_MIME_TYPE = "application/pdf"

TEXT_MIME_TYPE = "text/plain"

POWERPOINT_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.presentation"
)

EXCEL_XLSX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)

EXCEL_XLS_MIME_TYPE = "application/vnd.ms-excel"

DEFAULT_MAX_CHUNK_CHARACTERS = 2_000
DEFAULT_CHUNK_OVERLAP_CHARACTERS = 200


class FileExtractionError(RuntimeError):
    """Raised when a study file cannot be read safely."""


@dataclass(frozen=True)
class ExtractedSection:
    """One logical section extracted from a study file."""

    locator_type: str
    locator_label: str
    content: str

    metadata: dict[str, Any] = field(
        default_factory=dict,
    )


@dataclass(frozen=True)
class ExtractedDocument:
    """Complete normalized content extracted from a study file."""

    extracted_text: str
    sections: list[ExtractedSection]

    page_count: int | None = None
    slide_count: int | None = None
    sheet_count: int | None = None

    metadata: dict[str, Any] = field(
        default_factory=dict,
    )

    @property
    def character_count(self) -> int:
        """Return the number of extracted characters."""

        return len(
            self.extracted_text,
        )


@dataclass(frozen=True)
class ExtractedChunk:
    """One smaller section prepared for database storage."""

    chunk_index: int
    content: str

    locator_type: str
    locator_label: str

    token_count: int

    metadata: dict[str, Any] = field(
        default_factory=dict,
    )


def extract_document(
    payload: bytes,
    mime_type: str,
    filename: str,
) -> ExtractedDocument:
    """Extract readable content from a supported study file."""

    if not payload:
        raise FileExtractionError(
            "The uploaded study file is empty.",
        )

    normalized_mime_type = (
        mime_type.strip().lower()
    )

    if normalized_mime_type == PDF_MIME_TYPE:
        return extract_pdf(
            payload=payload,
            filename=filename,
        )

    if normalized_mime_type == TEXT_MIME_TYPE:
        return extract_plain_text(
            payload=payload,
            filename=filename,
        )

    if normalized_mime_type == POWERPOINT_MIME_TYPE:
        return extract_powerpoint(
            payload=payload,
            filename=filename,
        )

    if normalized_mime_type == EXCEL_XLSX_MIME_TYPE:
        return extract_excel_workbook(
            payload=payload,
            filename=filename,
        )

    if normalized_mime_type == EXCEL_XLS_MIME_TYPE:
        return extract_legacy_excel_workbook(
            payload=payload,
            filename=filename,
        )

    raise FileExtractionError(
        "This file type is not supported by the "
        "current extractor.",
    )


def extract_plain_text(
    payload: bytes,
    filename: str,
) -> ExtractedDocument:
    """Decode and normalize a UTF-8 plain-text document."""

    try:
        decoded_text = payload.decode(
            "utf-8-sig",
        )
    except UnicodeDecodeError as error:
        raise FileExtractionError(
            "The text file must use UTF-8 encoding.",
        ) from error

    normalized_text = normalize_extracted_text(
        decoded_text,
    )

    if not normalized_text:
        raise FileExtractionError(
            "The text file does not contain readable content.",
        )

    section = ExtractedSection(
        locator_type="document",
        locator_label="Complete document",
        content=normalized_text,
        metadata={
            "filename": filename,
        },
    )

    return ExtractedDocument(
        extracted_text=normalized_text,
        sections=[
            section,
        ],
        metadata={
            "filename": filename,
            "mime_type": TEXT_MIME_TYPE,
            "encoding": "utf-8",
        },
    )


def extract_pdf(
    payload: bytes,
    filename: str,
) -> ExtractedDocument:
    """Extract readable text from every PDF page."""

    try:
        reader = PdfReader(
            BytesIO(payload),
        )
    except (
        PdfReadError,
        ValueError,
        TypeError,
        OSError,
    ) as error:
        raise FileExtractionError(
            "The PDF file is damaged or unreadable.",
        ) from error

    if reader.is_encrypted:
        try:
            decrypt_result = reader.decrypt("")
        except Exception as error:
            raise FileExtractionError(
                "Password-protected PDFs are not supported.",
            ) from error

        if decrypt_result == 0:
            raise FileExtractionError(
                "Password-protected PDFs are not supported.",
            )

    page_count = len(
        reader.pages,
    )

    if page_count == 0:
        raise FileExtractionError(
            "The PDF does not contain any pages.",
        )

    sections: list[ExtractedSection] = []
    complete_page_texts: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        try:
            page_text = page.extract_text() or ""
        except Exception as error:
            raise FileExtractionError(
                f"Page {page_number} could not be read.",
            ) from error

        normalized_page_text = normalize_extracted_text(
            page_text,
        )

        if not normalized_page_text:
            continue

        sections.append(
            ExtractedSection(
                locator_type="page",
                locator_label=f"Page {page_number}",
                content=normalized_page_text,
                metadata={
                    "page_number": page_number,
                    "filename": filename,
                },
            ),
        )

        complete_page_texts.append(
            normalized_page_text,
        )

    extracted_text = normalize_extracted_text(
        "\n\n".join(
            complete_page_texts,
        ),
    )

    if not extracted_text:
        raise FileExtractionError(
            "No readable text was found in the PDF. "
            "The document may contain only scanned images.",
        )

    return ExtractedDocument(
        extracted_text=extracted_text,
        sections=sections,
        page_count=page_count,
        metadata={
            "filename": filename,
            "mime_type": PDF_MIME_TYPE,
            "pages_with_text": len(
                sections,
            ),
        },
    )


def extract_powerpoint(
    payload: bytes,
    filename: str,
) -> ExtractedDocument:
    """Extract text, tables, charts, and notes from a PPTX."""

    try:
        presentation = Presentation(
            BytesIO(payload),
        )
    except (
        PackageNotFoundError,
        PythonPptxError,
        BadZipFile,
        KeyError,
        ValueError,
        TypeError,
        OSError,
    ) as error:
        raise FileExtractionError(
            "The PowerPoint file is damaged or unreadable.",
        ) from error

    slide_count = len(
        presentation.slides,
    )

    if slide_count == 0:
        raise FileExtractionError(
            "The PowerPoint presentation has no slides.",
        )

    sections: list[ExtractedSection] = []
    complete_slide_texts: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_fragments: list[str] = []

        for shape in slide.shapes:
            slide_fragments.extend(
                extract_powerpoint_shape_text(
                    shape,
                ),
            )

        slide_body = normalize_extracted_text(
            "\n".join(
                slide_fragments,
            ),
        )

        notes_text = extract_powerpoint_notes(
            slide,
        )

        slide_parts: list[str] = []

        if slide_body:
            slide_parts.append(
                slide_body,
            )

        if notes_text:
            slide_parts.append(
                "Speaker notes:\n"
                f"{notes_text}",
            )

        slide_text = normalize_extracted_text(
            "\n\n".join(
                slide_parts,
            ),
        )

        if not slide_text:
            continue

        slide_title = get_powerpoint_slide_title(
            slide,
        )

        locator_label = (
            f"Slide {slide_number}: {slide_title}"
            if slide_title
            else f"Slide {slide_number}"
        )

        sections.append(
            ExtractedSection(
                locator_type="slide",
                locator_label=locator_label,
                content=slide_text,
                metadata={
                    "slide_number": slide_number,
                    "slide_title": slide_title or None,
                    "has_speaker_notes": bool(
                        notes_text,
                    ),
                    "filename": filename,
                },
            ),
        )

        complete_slide_texts.append(
            f"{locator_label}\n{slide_text}",
        )

    extracted_text = normalize_extracted_text(
        "\n\n".join(
            complete_slide_texts,
        ),
    )

    if not extracted_text:
        raise FileExtractionError(
            "No readable text was found in the "
            "PowerPoint presentation.",
        )

    return ExtractedDocument(
        extracted_text=extracted_text,
        sections=sections,
        slide_count=slide_count,
        metadata={
            "filename": filename,
            "mime_type": POWERPOINT_MIME_TYPE,
            "slides_with_text": len(
                sections,
            ),
        },
    )


def extract_powerpoint_shape_text(
    shape: Any,
) -> list[str]:
    """Extract readable text recursively from a slide shape."""

    fragments: list[str] = []

    if bool(
        getattr(
            shape,
            "has_text_frame",
            False,
        ),
    ):
        paragraph_texts = [
            normalize_extracted_text(
                paragraph.text,
            )
            for paragraph in shape.text_frame.paragraphs
        ]

        shape_text = normalize_extracted_text(
            "\n".join(
                paragraph_text
                for paragraph_text in paragraph_texts
                if paragraph_text
            ),
        )

        if shape_text:
            fragments.append(
                shape_text,
            )

    if bool(
        getattr(
            shape,
            "has_table",
            False,
        ),
    ):
        for row in shape.table.rows:
            cell_values = [
                normalize_extracted_text(
                    cell.text,
                )
                for cell in row.cells
            ]

            if any(
                cell_values,
            ):
                fragments.append(
                    " | ".join(
                        value or ""
                        for value in cell_values
                    ).strip(),
                )

    if bool(
        getattr(
            shape,
            "has_chart",
            False,
        ),
    ):
        chart = shape.chart

        if bool(
            getattr(
                chart,
                "has_title",
                False,
            ),
        ):
            chart_title = chart.chart_title

            if bool(
                getattr(
                    chart_title,
                    "has_text_frame",
                    False,
                ),
            ):
                title_text = normalize_extracted_text(
                    chart_title.text_frame.text,
                )

                if title_text:
                    fragments.append(
                        f"Chart: {title_text}",
                    )

    child_shapes = getattr(
        shape,
        "shapes",
        None,
    )

    if child_shapes is not None:
        for child_shape in child_shapes:
            fragments.extend(
                extract_powerpoint_shape_text(
                    child_shape,
                ),
            )

    return fragments


def extract_powerpoint_notes(
    slide: Any,
) -> str:
    """Return speaker notes belonging to one slide."""

    try:
        if not slide.has_notes_slide:
            return ""

        notes_text_frame = (
            slide.notes_slide.notes_text_frame
        )

        if notes_text_frame is None:
            return ""

        note_lines = [
            normalize_extracted_text(
                paragraph.text,
            )
            for paragraph in notes_text_frame.paragraphs
        ]

        return normalize_extracted_text(
            "\n".join(
                line
                for line in note_lines
                if line
            ),
        )

    except (
        AttributeError,
        ValueError,
    ):
        return ""


def get_powerpoint_slide_title(
    slide: Any,
) -> str:
    """Return a normalized slide title when available."""

    title_shape = getattr(
        slide.shapes,
        "title",
        None,
    )

    if title_shape is None:
        return ""

    return normalize_extracted_text(
        str(
            getattr(
                title_shape,
                "text",
                "",
            ),
        ),
    )


def extract_excel_workbook(
    payload: bytes,
    filename: str,
) -> ExtractedDocument:
    """Extract readable cell values from an XLSX workbook."""

    try:
        workbook = load_workbook(
            filename=BytesIO(payload),
            read_only=True,
            data_only=True,
        )
    except (
        InvalidFileException,
        BadZipFile,
        KeyError,
        ValueError,
        TypeError,
        OSError,
    ) as error:
        raise FileExtractionError(
            "The Excel workbook is damaged or unreadable.",
        ) from error

    sheet_count = len(
        workbook.worksheets,
    )

    if sheet_count == 0:
        workbook.close()

        raise FileExtractionError(
            "The Excel workbook does not contain "
            "any worksheets.",
        )

    sections: list[ExtractedSection] = []
    complete_sheet_texts: list[str] = []

    try:
        for sheet_number, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):
            row_fragments: list[str] = []

            populated_row_count = 0
            populated_cell_count = 0

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    values_only=True,
                ),
                start=1,
            ):
                formatted_values = [
                    normalize_excel_cell_value(
                        value,
                    )
                    for value in row
                ]

                while (
                    formatted_values
                    and not formatted_values[-1]
                ):
                    formatted_values.pop()

                if not any(
                    formatted_values,
                ):
                    continue

                populated_row_count += 1

                populated_cell_count += sum(
                    1
                    for value in formatted_values
                    if value
                )

                row_fragments.append(
                    f"Row {row_number}: "
                    + " | ".join(
                        formatted_values,
                    ),
                )

            sheet_text = normalize_extracted_text(
                "\n".join(
                    row_fragments,
                ),
            )

            if not sheet_text:
                continue

            locator_label = (
                f"Sheet {sheet_number}: "
                f"{worksheet.title}"
            )

            sections.append(
                ExtractedSection(
                    locator_type="sheet",
                    locator_label=locator_label,
                    content=sheet_text,
                    metadata={
                        "sheet_number": sheet_number,
                        "sheet_name": worksheet.title,
                        "populated_row_count": (
                            populated_row_count
                        ),
                        "populated_cell_count": (
                            populated_cell_count
                        ),
                        "filename": filename,
                        "workbook_format": "xlsx",
                    },
                ),
            )

            complete_sheet_texts.append(
                f"{locator_label}\n{sheet_text}",
            )

    finally:
        workbook.close()

    extracted_text = normalize_extracted_text(
        "\n\n".join(
            complete_sheet_texts,
        ),
    )

    if not extracted_text:
        raise FileExtractionError(
            "No readable cell values were found in "
            "the Excel workbook.",
        )

    return ExtractedDocument(
        extracted_text=extracted_text,
        sections=sections,
        sheet_count=sheet_count,
        metadata={
            "filename": filename,
            "mime_type": EXCEL_XLSX_MIME_TYPE,
            "workbook_format": "xlsx",
            "sheets_with_content": len(
                sections,
            ),
        },
    )


def extract_legacy_excel_workbook(
    payload: bytes,
    filename: str,
) -> ExtractedDocument:
    """Extract readable values from a legacy XLS workbook."""

    try:
        workbook = xlrd.open_workbook(
            file_contents=payload,
            on_demand=True,
            ragged_rows=True,
        )
    except Exception as error:
        raise FileExtractionError(
            "The legacy Excel workbook is damaged or unreadable.",
        ) from error

    sheet_count = workbook.nsheets

    if sheet_count == 0:
        workbook.release_resources()

        raise FileExtractionError(
            "The legacy Excel workbook does not contain "
            "any worksheets.",
        )

    sections: list[ExtractedSection] = []
    complete_sheet_texts: list[str] = []

    try:
        for sheet_index in range(
            sheet_count,
        ):
            worksheet = workbook.sheet_by_index(
                sheet_index,
            )

            sheet_number = sheet_index + 1

            row_fragments: list[str] = []

            populated_row_count = 0
            populated_cell_count = 0

            for row_index in range(
                worksheet.nrows,
            ):
                row_length = worksheet.row_len(
                    row_index,
                )

                formatted_values = [
                    normalize_legacy_excel_cell(
                        cell=worksheet.cell(
                            row_index,
                            column_index,
                        ),
                        datemode=workbook.datemode,
                    )
                    for column_index in range(
                        row_length,
                    )
                ]

                while (
                    formatted_values
                    and not formatted_values[-1]
                ):
                    formatted_values.pop()

                if not any(
                    formatted_values,
                ):
                    continue

                populated_row_count += 1

                populated_cell_count += sum(
                    1
                    for value in formatted_values
                    if value
                )

                row_fragments.append(
                    f"Row {row_index + 1}: "
                    + " | ".join(
                        formatted_values,
                    ),
                )

            sheet_text = normalize_extracted_text(
                "\n".join(
                    row_fragments,
                ),
            )

            if not sheet_text:
                continue

            locator_label = (
                f"Sheet {sheet_number}: "
                f"{worksheet.name}"
            )

            sections.append(
                ExtractedSection(
                    locator_type="sheet",
                    locator_label=locator_label,
                    content=sheet_text,
                    metadata={
                        "sheet_number": sheet_number,
                        "sheet_name": worksheet.name,
                        "populated_row_count": (
                            populated_row_count
                        ),
                        "populated_cell_count": (
                            populated_cell_count
                        ),
                        "filename": filename,
                        "workbook_format": "xls",
                    },
                ),
            )

            complete_sheet_texts.append(
                f"{locator_label}\n{sheet_text}",
            )

    except FileExtractionError:
        raise

    except Exception as error:
        raise FileExtractionError(
            "The legacy Excel workbook could not be read.",
        ) from error

    finally:
        workbook.release_resources()

    extracted_text = normalize_extracted_text(
        "\n\n".join(
            complete_sheet_texts,
        ),
    )

    if not extracted_text:
        raise FileExtractionError(
            "No readable cell values were found in "
            "the legacy Excel workbook.",
        )

    return ExtractedDocument(
        extracted_text=extracted_text,
        sections=sections,
        sheet_count=sheet_count,
        metadata={
            "filename": filename,
            "mime_type": EXCEL_XLS_MIME_TYPE,
            "workbook_format": "xls",
            "sheets_with_content": len(
                sections,
            ),
        },
    )


def normalize_excel_cell_value(
    value: Any,
) -> str:
    """Convert an XLSX cell value into readable text."""

    if value is None:
        return ""

    if isinstance(
        value,
        bool,
    ):
        return "TRUE" if value else "FALSE"

    if isinstance(
        value,
        datetime,
    ):
        return value.isoformat(
            sep=" ",
        )

    if isinstance(
        value,
        (
            date,
            time,
        ),
    ):
        return value.isoformat()

    if isinstance(
        value,
        Decimal,
    ):
        return format(
            value,
            "f",
        )

    if isinstance(
        value,
        float,
    ):
        return normalize_numeric_value(
            value,
        )

    return normalize_extracted_text(
        str(value),
    )


def normalize_legacy_excel_cell(
    cell: Any,
    datemode: int,
) -> str:
    """Convert an xlrd legacy XLS cell into readable text."""

    if cell.ctype in {
        xlrd.XL_CELL_EMPTY,
        xlrd.XL_CELL_BLANK,
    }:
        return ""

    if cell.ctype == xlrd.XL_CELL_TEXT:
        return normalize_extracted_text(
            str(
                cell.value,
            ),
        )

    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return normalize_numeric_value(
            float(
                cell.value,
            ),
        )

    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return (
            "TRUE"
            if bool(
                cell.value,
            )
            else "FALSE"
        )

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            converted_datetime = xldate_as_datetime(
                cell.value,
                datemode,
            )
        except (
            XLDateError,
            ValueError,
            OverflowError,
        ):
            return normalize_numeric_value(
                float(
                    cell.value,
                ),
            )

        if converted_datetime.time() == time.min:
            return converted_datetime.date().isoformat()

        return converted_datetime.isoformat(
            sep=" ",
        )

    if cell.ctype == xlrd.XL_CELL_ERROR:
        try:
            error_code = int(
                cell.value,
            )
        except (
            TypeError,
            ValueError,
        ):
            return "#ERROR"

        error_text = xlrd.error_text_from_code.get(
            error_code,
            "Unknown error",
        )

        return f"#ERROR: {error_text}"

    return normalize_extracted_text(
        str(
            cell.value,
        ),
    )


def normalize_numeric_value(
    value: float,
) -> str:
    """Convert a spreadsheet number into compact text."""

    if value.is_integer():
        return str(
            int(
                value,
            ),
        )

    return format(
        value,
        ".15g",
    )


def chunk_extracted_document(
    document: ExtractedDocument,
    max_characters: int = (
        DEFAULT_MAX_CHUNK_CHARACTERS
    ),
    overlap_characters: int = (
        DEFAULT_CHUNK_OVERLAP_CHARACTERS
    ),
) -> list[ExtractedChunk]:
    """Split extracted sections into overlapping text chunks."""

    if max_characters <= 0:
        raise ValueError(
            "max_characters must be greater than zero.",
        )

    if overlap_characters < 0:
        raise ValueError(
            "overlap_characters cannot be negative.",
        )

    if overlap_characters >= max_characters:
        raise ValueError(
            "overlap_characters must be smaller than "
            "max_characters.",
        )

    chunks: list[ExtractedChunk] = []

    for section in document.sections:
        section_chunks = split_text_into_chunks(
            text=section.content,
            max_characters=max_characters,
            overlap_characters=overlap_characters,
        )

        for section_chunk_index, content in enumerate(
            section_chunks,
        ):
            chunks.append(
                ExtractedChunk(
                    chunk_index=len(
                        chunks,
                    ),
                    content=content,
                    locator_type=section.locator_type,
                    locator_label=section.locator_label,
                    token_count=estimate_token_count(
                        content,
                    ),
                    metadata={
                        **section.metadata,
                        "section_chunk_index": (
                            section_chunk_index
                        ),
                    },
                ),
            )

    if not chunks:
        raise FileExtractionError(
            "The extracted document did not produce "
            "any readable chunks.",
        )

    return chunks


def split_text_into_chunks(
    text: str,
    max_characters: int,
    overlap_characters: int,
) -> list[str]:
    """Split text while preferring natural boundaries."""

    normalized_text = normalize_extracted_text(
        text,
    )

    if not normalized_text:
        return []

    chunks: list[str] = []

    start_index = 0
    text_length = len(
        normalized_text,
    )

    while start_index < text_length:
        proposed_end = min(
            start_index + max_characters,
            text_length,
        )

        end_index = proposed_end

        if proposed_end < text_length:
            end_index = find_chunk_boundary(
                text=normalized_text,
                start_index=start_index,
                proposed_end=proposed_end,
            )

        chunk = normalized_text[
            start_index:end_index
        ].strip()

        if chunk:
            chunks.append(
                chunk,
            )

        if end_index >= text_length:
            break

        next_start = max(
            end_index - overlap_characters,
            start_index + 1,
        )

        while (
            next_start < text_length
            and normalized_text[next_start].isspace()
        ):
            next_start += 1

        start_index = next_start

    return chunks


def find_chunk_boundary(
    text: str,
    start_index: int,
    proposed_end: int,
) -> int:
    """Find a natural break before the maximum size."""

    minimum_boundary = (
        start_index
        + int(
            (
                proposed_end
                - start_index
            )
            * 0.6,
        )
    )

    boundary_candidates = [
        text.rfind(
            "\n\n",
            minimum_boundary,
            proposed_end,
        ),
        text.rfind(
            ". ",
            minimum_boundary,
            proposed_end,
        ),
        text.rfind(
            " ",
            minimum_boundary,
            proposed_end,
        ),
    ]

    best_boundary = max(
        boundary_candidates,
    )

    if best_boundary <= start_index:
        return proposed_end

    if text.startswith(
        "\n\n",
        best_boundary,
    ):
        return best_boundary

    return best_boundary + 1


def normalize_extracted_text(
    text: str,
) -> str:
    """Normalize whitespace without removing paragraphs."""

    normalized = (
        text.replace(
            "\x00",
            "",
        )
        .replace(
            "\v",
            "\n",
        )
        .replace(
            "\r\n",
            "\n",
        )
        .replace(
            "\r",
            "\n",
        )
    )

    normalized_lines = [
        re.sub(
            r"[ \t]+",
            " ",
            line,
        ).strip()
        for line in normalized.split(
            "\n",
        )
    ]

    normalized = "\n".join(
        normalized_lines,
    )

    normalized = re.sub(
        r"\n{3,}",
        "\n\n",
        normalized,
    )

    return normalized.strip()


def estimate_token_count(
    text: str,
) -> int:
    """Return a lightweight token estimate for metadata."""

    word_count = len(
        re.findall(
            r"\S+",
            text,
        ),
    )

    if word_count == 0:
        return 0

    return max(
        1,
        round(
            word_count * 1.3,
        ),
    )